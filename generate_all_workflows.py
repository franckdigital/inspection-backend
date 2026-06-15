# -*- coding: utf-8 -*-
"""
Génère : Workflows_Complets_eInspection_CI_2026.docx
Documentation exhaustive de tous les modules de la Plateforme e-Inspection du Travail
Ministère de l'Emploi et de la Protection Sociale — Côte d'Ivoire
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import date

# ─── Palette Côte d'Ivoire ───────────────────────────────────────────────────
ORANGE  = RGBColor(0xF4, 0x7E, 0x1C)
VERT    = RGBColor(0x00, 0x9A, 0x44)
ROUGE   = RGBColor(0xDC, 0x26, 0x26)
BLEU    = RGBColor(0x1D, 0x4E, 0xD8)
VIOLET  = RGBColor(0x7C, 0x3A, 0xED)
TEAL    = RGBColor(0x0F, 0x76, 0x6E)
GRIS    = RGBColor(0x6B, 0x72, 0x80)
NOIR    = RGBColor(0x11, 0x18, 0x27)
BLANC   = RGBColor(0xFF, 0xFF, 0xFF)

HEX = {
    "orange": "F47E1C", "vert": "009A44", "rouge": "DC2626",
    "bleu": "1D4ED8", "violet": "7C3AED", "teal": "0F766E",
    "gris": "6B7280", "noir": "111827", "blanc": "FFFFFF",
    "bg_orange": "FFF7ED", "bg_vert": "F0FDF4", "bg_bleu": "EFF6FF",
    "bg_rouge": "FEF2F2", "bg_violet": "F5F3FF", "bg_teal": "F0FDFA",
    "bg_gris": "F9FAFB", "bg_jaune": "FEFCE8",
}

doc = Document()

# ─── Marges ──────────────────────────────────────────────────────────────────
for sec in doc.sections:
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.0)


# ─── Utilitaires ─────────────────────────────────────────────────────────────
def cell_bg(cell, hex_color):
    tc = cell._tc
    pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pr.append(shd)


def cell_borders(cell, color="E5E7EB"):
    tc = cell._tc
    pr = tc.get_or_add_tcPr()
    bdr = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), color)
        bdr.append(b)
    pr.append(bdr)


def h1(text, color=ORANGE, bar_hex="F47E1C"):
    p = doc.add_paragraph()
    p.clear()
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = color
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(5)
    bdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "8")
    bot.set(qn("w:space"), "2")
    bot.set(qn("w:color"), bar_hex)
    bdr.append(bot)
    p._p.get_or_add_pPr().append(bdr)
    return p


def h2(text, color=BLEU):
    p = doc.add_paragraph()
    p.clear()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = color
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    return p


def h3(text, color=NOIR):
    p = doc.add_paragraph()
    p.clear()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = color
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    return p


def body(text, italic=False, color=None, size=10.5):
    p = doc.add_paragraph()
    p.clear()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    if color:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    return p


def bul(text, lvl=0, color=None):
    p = doc.add_paragraph()
    p.clear()
    prefix = "    " * lvl + ("• " if lvl == 0 else "◦ ")
    r = p.add_run(prefix + text)
    r.font.size = Pt(10)
    if color:
        r.font.color.rgb = color
    p.paragraph_format.left_indent = Cm(0.4 * (lvl + 1))
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    return p


def tbl(headers, rows, hdr_bg="009A44", col_widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    # En-tête
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        cell_bg(c, hdr_bg)
        pp = c.paragraphs[0]
        pp.clear()
        rr = pp.add_run(h)
        rr.bold = True
        rr.font.size = Pt(9)
        rr.font.color.rgb = BLANC
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Données
    for ri, row_data in enumerate(rows):
        bg = "F9FAFB" if ri % 2 == 0 else "FFFFFF"
        for ci, val in enumerate(row_data):
            c = t.rows[ri + 1].cells[ci]
            cell_bg(c, bg)
            cell_borders(c)
            pp = c.paragraphs[0]
            pp.clear()
            rr = pp.add_run(str(val))
            rr.font.size = Pt(9)
    doc.add_paragraph()
    return t


def flow(steps, bg_hex="009A44"):
    t = doc.add_table(rows=2, cols=len(steps))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (label, desc) in enumerate(steps):
        c1 = t.rows[0].cells[i]
        cell_bg(c1, bg_hex)
        pp = c1.paragraphs[0]
        pp.clear()
        r = pp.add_run(label)
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = BLANC
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER

        c2 = t.rows[1].cells[i]
        cell_bg(c2, "F8FAFC")
        pp = c2.paragraphs[0]
        pp.clear()
        r = pp.add_run(desc)
        r.font.size = Pt(8)
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    return t


def badge_row(items):
    """Affiche des statuts sur une ligne (table 1xN)."""
    colors = ["1D4ED8", "F47E1C", "009A44", "DC2626", "7C3AED", "0F766E", "6B7280", "92400E"]
    t = doc.add_table(rows=1, cols=len(items))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, item in enumerate(items):
        c = t.rows[0].cells[i]
        cell_bg(c, colors[i % len(colors)])
        pp = c.paragraphs[0]
        pp.clear()
        r = pp.add_run(item)
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = BLANC
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


def info_box(text, bg="FFF7ED", border="F47E1C"):
    p = doc.add_paragraph()
    p.clear()
    r = p.add_run(text)
    r.font.size = Pt(9.5)
    r.italic = True
    r.font.color.rgb = RGBColor(
        int(border[0:2], 16), int(border[2:4], 16), int(border[4:6], 16)
    )
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    return p


# ════════════════════════════════════════════════════════════════════════════
#  COUVERTURE
# ════════════════════════════════════════════════════════════════════════════
def cover():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(72)
    r = p.add_run("DOCUMENTATION COMPLÈTE DES WORKFLOWS")
    r.bold = True; r.font.size = Pt(26); r.font.color.rgb = ORANGE

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Plateforme Nationale e-Inspection du Travail")
    r.bold = True; r.font.size = Pt(17); r.font.color.rgb = VERT

    doc.add_paragraph()

    for line, sz, col, ital in [
        ("Ministère de l'Emploi et de la Protection Sociale", 13, GRIS, False),
        ("Direction Générale du Travail (DGT)", 12, GRIS, True),
        ("République de Côte d'Ivoire", 12, GRIS, True),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.font.size = Pt(sz); r.font.color.rgb = col; r.italic = ital

    doc.add_paragraph(); doc.add_paragraph()

    meta = [
        ("Référence",  "DOC-WF-ALL-EINSP-CI-2026-001"),
        ("Version",    "2.0 — Complète tous modules"),
        ("Date",       date.today().strftime("%d %B %Y")),
        ("Modules",    "15 applications Django — Backend, Frontend, Mobile"),
        ("Diffusion",  "Usage interne — Confidentiel"),
    ]
    t = doc.add_table(rows=len(meta), cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(meta):
        bg = "F47E1C" if i == 0 else ("F9FAFB" if i % 2 == 0 else "FFFFFF")
        c0, c1 = t.rows[i].cells[0], t.rows[i].cells[1]
        cell_bg(c0, bg); cell_bg(c1, "FFFFFF" if i > 0 else "FFF7ED")
        pp = c0.paragraphs[0]; pp.clear()
        r = pp.add_run(k); r.bold = True; r.font.size = Pt(9.5)
        r.font.color.rgb = BLANC if i == 0 else NOIR
        pp = c1.paragraphs[0]; pp.clear()
        r = pp.add_run(v); r.font.size = Pt(9.5)

    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  SOMMAIRE
# ════════════════════════════════════════════════════════════════════════════
def sommaire():
    h1("Sommaire")
    sections = [
        ("1.",  "Architecture globale de la plateforme"),
        ("2.",  "Rôles, permissions et RBAC"),
        ("3.",  "Module USERS — Authentification et profils"),
        ("4.",  "Module ENTERPRISES — Gestion des entreprises"),
        ("5.",  "Module COMPLAINTS — Gestion des plaintes"),
        ("6.",  "Module INSPECTIONS — Contrôles terrain et zones"),
        ("7.",  "Module MEDIATIONS — Médiation et conciliation"),
        ("8.",  "Module HIERARCHY — Escalade et approbations"),
        ("9.",  "Module JUDICIAL — Procédures judiciaires"),
        ("10.", "Module DOMESTIC — Travailleurs domestiques"),
        ("11.", "Module AI — Intelligence artificielle"),
        ("12.", "Module OBSERVATORY — Statistiques nationales"),
        ("13.", "Module NOTIFICATIONS — Emails et SMS"),
        ("14.", "Module GED — Gestion électronique de documents"),
        ("15.", "Module ADMINISTRATION — Configuration système"),
        ("16.", "Module BI — Business Intelligence"),
        ("17.", "Module LANDING — Site web public"),
        ("18.", "Tâches Celery automatiques"),
        ("19.", "Référentiel API complet"),
        ("20.", "Architecture Frontend (React / TypeScript)"),
        ("21.", "Architecture Mobile (React Native / Expo)"),
    ]
    for num, title in sections:
        p = doc.add_paragraph()
        p.clear()
        r = p.add_run(f"{num:<6}{title}")
        r.font.size = Pt(10.5)
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(2)
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  1. ARCHITECTURE GLOBALE
# ════════════════════════════════════════════════════════════════════════════
def section_architecture():
    h1("1. Architecture globale de la plateforme")

    body(
        "La Plateforme Nationale e-Inspection du Travail est un système d'information "
        "multi-couches déployé par la Direction Générale du Travail (DGT) de Côte d'Ivoire. "
        "Elle couvre le cycle complet de traitement des conflits du travail, depuis le dépôt "
        "d'une plainte jusqu'à la décision judiciaire, en passant par la médiation, "
        "conformément à la Loi n° 2015-532 du 20 juillet 2015."
    )

    h2("Stack technique")
    tbl(
        ["Couche", "Technologie", "Version", "Rôle"],
        [
            ["API Backend",       "Django + Django REST Framework", "5.0.6 / 3.15.1", "Logique métier, API REST, auth"],
            ["Base de données",   "MySQL / MariaDB",                "8.0+",            "Persistance relationnelle"],
            ["Cache & Files",     "Redis",                          "7+",              "Queue Celery, sessions, cache"],
            ["Tâches async",      "Celery",                         "5.3.6",           "Emails, SMS, relances planifiées"],
            ["Temps réel",        "Django Channels",                "4.1.0",           "WebSocket, notifications push"],
            ["IA générative",     "OpenAI GPT-4 / Whisper",         "API",             "Chatbot légal, transcription vocale"],
            ["Frontend web",      "React + TypeScript",             "18+",             "SPA pour inspecteurs & employeurs"],
            ["Mobile",            "React Native + Expo",            "50+",             "iOS et Android"],
            ["Documents",         "python-docx",                    "1.2.0",           "PV, accords, rapports Word"],
            ["Auth",              "SimpleJWT + pyotp",              "5.3.1",           "JWT access/refresh + TOTP 2FA"],
            ["Timezone",          "Africa/Abidjan",                 "UTC+0",           "Fuseau horaire Côte d'Ivoire"],
            ["SMS",               "Orange CI / MTN CI / Moov",      "API",             "Convocations et alertes SMS (+225)"],
        ],
        hdr_bg="F47E1C",
    )

    h2("15 applications Django")
    tbl(
        ["App", "URL prefix", "Rôle principal"],
        [
            ["users",          "/api/users/",          "Authentification, profils, 2FA, RBAC"],
            ["enterprises",    "/api/enterprises/",    "Entreprises, branches, conformité"],
            ["complaints",     "/api/complaints/",     "Plaintes, workflow de traitement"],
            ["inspections",    "/api/inspections/",    "Zones, contrôles terrain, affectations"],
            ["mediations",     "/api/mediations/",     "Séances de médiation, convocations, accords"],
            ["hierarchy",      "/api/hierarchy/",      "Escalade, approbations, délégations"],
            ["judicial",       "/api/judicial/",       "Procédures judiciaires, audiences, décisions"],
            ["domestic",       "/api/domestic/",       "Travailleurs domestiques, contrats, pointage"],
            ["ai",             "/api/ai/",             "Chatbot légal, analyse documentaire, IA"],
            ["observatory",    "/api/observatory/",    "Statistiques nationales, indicateurs"],
            ["notifications",  "/api/notifications/",  "Templates email/SMS, logs d'envoi"],
            ["ged",            "/api/ged/",            "Gestion électronique de documents"],
            ["administration", "/api/admin/",          "Configuration, audit, maintenance"],
            ["bi",             "/api/bi/",             "Dashboards BI, KPIs, rapports"],
            ["landing",        "/api/public/",         "Site public, actualités, FAQ, contact"],
        ],
        hdr_bg="009A44",
    )
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  2. RBAC
# ════════════════════════════════════════════════════════════════════════════
def section_rbac():
    h1("2. Rôles, permissions et RBAC")

    body(
        "Le système implémente un contrôle d'accès basé sur les rôles (RBAC) "
        "avec 8 profils distincts. Chaque rôle détermine les modules accessibles, "
        "les actions autorisées et le périmètre géographique."
    )

    tbl(
        ["Rôle", "Code", "Périmètre", "Accès clés"],
        [
            ["Travailleur (salarié)",    "EMPLOYE",           "Ses propres dossiers",      "Dépôt plainte, suivi, médiation en lecture"],
            ["Travailleur domestique",   "EMPLOYE_MAISON",    "Ses propres dossiers",      "Pointage, congés, bulletins, plaintes vocales"],
            ["Employeur",                "EMPLOYEUR",         "Son entreprise",            "Gestion employés, validation heures, convocations"],
            ["Inspecteur du Travail",    "INSPECTEUR",        "Plaintes assignées",        "CRUD dossiers, contrôles, médiations, PV"],
            ["Chef d'inspection",        "CHEF_INSPECTION",   "Son inspection",            "Supervision, assignation, escalade, validation"],
            ["Directeur Régional",       "DIRECTEUR_REGIONAL","Sa région",                 "Tableau de bord régional, rapports"],
            ["Directeur Général",        "DIRECTEUR_GENERAL", "National",                  "Vue globale, validation hiérarchique"],
            ["Administrateur système",   "ADMIN",             "Tout le système",           "CRUD illimité, configuration, audit"],
        ],
        hdr_bg="7C3AED",
    )

    h2("Matrice de permissions par module")
    tbl(
        ["Module", "EMPLOYE", "EMPLOYEUR", "INSPECTEUR", "CHEF", "DIR.REG.", "DIR.GEN.", "ADMIN"],
        [
            ["Users",         "R (soi)", "R (soi)",  "R",         "R",    "R",    "R",    "CRUD"],
            ["Enterprises",   "R",       "R/U",      "R/U",       "CRUD", "CRUD", "CRUD", "CRUD"],
            ["Complaints",    "CRU",     "R",        "CRUD",      "CRUD", "R",    "R",    "CRUD"],
            ["Inspections",   "-",       "R",        "CRUD",      "CRUD", "R",    "R",    "CRUD"],
            ["Mediations",    "R",       "R",        "CRUD",      "CRUD", "R",    "R",    "CRUD"],
            ["Hierarchy",     "-",       "-",        "R",         "CRUD", "CRUD", "CRUD", "CRUD"],
            ["Judicial",      "R",       "R",        "CRUD",      "CRUD", "R",    "R",    "CRUD"],
            ["Domestic",      "CRUD*",   "CRUD*",    "CRUD",      "CRUD", "R",    "R",    "CRUD"],
            ["AI",            "R",       "R",        "CRUD",      "R",    "R",    "R",    "CRUD"],
            ["Observatory",   "-",       "-",        "R",         "R",    "R",    "CRUD", "CRUD"],
            ["Notifications", "-",       "-",        "R",         "R",    "R",    "R",    "CRUD"],
            ["GED",           "-",       "-",        "CRUD",      "CRUD", "R",    "R",    "CRUD"],
            ["Administration","-",       "-",        "-",         "R",    "R",    "R",    "CRUD"],
            ["BI",            "-",       "-",        "R",         "CRUD", "CRUD", "CRUD", "CRUD"],
            ["Landing",       "R",       "R",        "R",         "R",    "R",    "R",    "CRUD"],
        ],
        hdr_bg="111827",
    )
    body("* Périmètre restreint aux propres données de l'utilisateur", italic=True, color=GRIS)
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  3. MODULE USERS
# ════════════════════════════════════════════════════════════════════════════
def section_users():
    h1("3. Module USERS — Authentification et profils", color=BLEU, bar_hex="1D4ED8")

    body("Gère l'inscription, l'authentification JWT, la vérification email, "
         "la 2FA (TOTP), la réinitialisation de mot de passe, et les profils spécialisés.")

    h2("Modèles")
    tbl(
        ["Modèle", "Champs clés", "Rôle"],
        [
            ["User",              "email, user_type, is_active, is_verified, otp_enabled",       "Compte utilisateur central"],
            ["EmployeeProfile",   "job_title, employer (FK Enterprise), assigned_inspector",     "Profil travailleur salarié"],
            ["InspectorProfile",  "badge_number, inspection_zone (FK Zone), specialization",     "Profil inspecteur"],
            ["EmployerProfile",   "employer_type (ENTERPRISE/AGENCY/INDIVIDUAL)",                "Profil employeur"],
            ["PasswordReset",     "token, expires_at, is_used",                                  "Réinitialisation mot de passe"],
            ["EmailVerification", "token, expires_at",                                           "Vérification adresse email"],
        ],
        hdr_bg="1D4ED8",
    )

    h2("Workflow d'inscription et vérification")
    flow([
        ("Inscription\nPOST /register/",    "Création compte\nis_active=False"),
        ("Email vérif\nenvoyé",             "Token 24h\nenvoyé par email"),
        ("Clic lien\n/verify-email/",       "is_verified=True\nis_active=True"),
        ("Connexion\nPOST /login/",         "JWT access+\nrefresh tokens"),
        ("2FA (optionnel)\n/otp/verify/",   "TOTP 6 chiffres\nrequis si activé"),
    ], bg_hex="1D4ED8")

    h2("Endpoints clés")
    tbl(
        ["Méthode", "Endpoint", "Description", "Auth"],
        [
            ["POST", "/api/users/register/",              "Inscription",                          "Non"],
            ["POST", "/api/users/login/",                 "Connexion JWT",                        "Non"],
            ["POST", "/api/users/token/refresh/",         "Rafraîchir le token",                  "Non"],
            ["GET/PATCH", "/api/users/me/",               "Profil courant",                       "JWT"],
            ["POST", "/api/users/change-password/",       "Changer mot de passe",                 "JWT"],
            ["POST", "/api/users/password-reset/",        "Demande réinitialisation",             "Non"],
            ["POST", "/api/users/password-reset/confirm/","Confirmer réinitialisation",           "Non"],
            ["POST", "/api/users/verify-email/",          "Vérifier email (token)",               "Non"],
            ["POST", "/api/users/otp/setup/",             "Activer 2FA",                          "JWT"],
            ["POST", "/api/users/otp/verify/",            "Vérifier code TOTP",                   "JWT"],
            ["GET/PATCH", "/api/users/profile/employee/", "Profil employé",                       "JWT"],
            ["GET/PATCH", "/api/users/profile/inspector/","Profil inspecteur",                    "JWT"],
            ["GET/PATCH", "/api/users/profile/employer/", "Profil employeur",                     "JWT"],
            ["GET", "/api/users/employees/unassigned/",   "Employés sans inspecteur",             "JWT CHEF+"],
            ["POST", "/api/users/employees/{id}/assign-inspector/", "Assigner inspecteur",       "JWT CHEF+"],
            ["POST", "/api/users/employees/bulk-auto-assign/",      "Auto-assignation masse",    "JWT ADMIN"],
        ],
        hdr_bg="1D4ED8",
    )

    h2("Tokens JWT")
    tbl(
        ["Token", "Durée", "Usage"],
        [
            ["Access Token",   "60 minutes",  "Inclus dans Authorization: Bearer — toutes les requêtes API"],
            ["Refresh Token",  "7 jours",     "Renouvelle l'access token sans re-connexion"],
            ["OTP TOTP",       "30 secondes", "Code 6 chiffres (Google Authenticator compatible)"],
            ["Email verify",   "24 heures",   "Lien de vérification email envoyé à l'inscription"],
            ["Password reset", "1 heure",     "Lien de réinitialisation envoyé par email"],
        ],
        hdr_bg="7C3AED",
    )
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  4. MODULE ENTERPRISES
# ════════════════════════════════════════════════════════════════════════════
def section_enterprises():
    h1("4. Module ENTERPRISES — Gestion des entreprises", color=TEAL, bar_hex="0F766E")

    body("Gère le référentiel des entreprises, leurs branches, leurs documents "
         "administratifs, et leur score de conformité. Alimenté lors des contrôles "
         "et mis à jour automatiquement lors des infractions constatées.")

    h2("Modèles")
    tbl(
        ["Modèle", "Champs clés", "Rôle"],
        [
            ["Enterprise",       "name, rccm, nif, sector, risk_level, compliance_score, is_verified",  "Entreprise principale"],
            ["EnterpriseBranch", "enterprise (FK), name, address, city, commune, manager",               "Succursale / filiale"],
            ["EnterpriseDocument","enterprise, doc_type, file, is_verified",                             "Documents légaux (RCCM, NIF, CNPS…)"],
            ["EnterpriseHistory","enterprise, event_type, description, actor",                           "Journal d'audit de l'entreprise"],
        ],
        hdr_bg="0F766E",
    )

    h2("Niveaux de risque et score de conformité")
    tbl(
        ["Niveau", "Code", "Score", "Fréquence inspection", "Actions automatiques"],
        [
            ["Faible",    "LOW",      "80–100", "Annuelle",        "Aucune alerte"],
            ["Moyen",     "MEDIUM",   "60–79",  "Semestrielle",    "Notification chef d'inspection"],
            ["Élevé",     "HIGH",     "40–59",  "Trimestrielle",   "Alerte Directeur Régional"],
            ["Critique",  "CRITICAL", "0–39",   "Mensuelle",       "Escalade DGT + signalement parquet"],
        ],
        hdr_bg="DC2626",
    )

    h2("Événements EnterpriseHistory")
    badge_row(["CREATION", "UPDATE", "INSPECTION", "COMPLAINT", "SANCTION", "PROCEDURE"])

    h2("Impact sur le score de conformité")
    tbl(
        ["Événement", "Impact score"],
        [
            ["Non-comparution à une médiation",          "−15 points"],
            ["Accord de médiation rompu",                "−20 points"],
            ["Infraction constatée en contrôle",         "−10 à −30 points (selon gravité)"],
            ["Contrôle CONFORME",                        "+10 points"],
            ["Documents légaux vérifiés et à jour",      "+5 points"],
        ],
        hdr_bg="F47E1C",
    )
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  5. MODULE COMPLAINTS
# ════════════════════════════════════════════════════════════════════════════
def section_complaints():
    h1("5. Module COMPLAINTS — Gestion des plaintes", color=ROUGE, bar_hex="DC2626")

    body("Cœur du système. Toute plainte déposée par un travailleur suit un workflow "
         "multi-étapes supervisé par l'Inspection du Travail, jusqu'à la résolution "
         "amiable ou judiciaire.")

    h2("Workflow principal des statuts")
    flow([
        ("PENDING\n(En attente)",     "Plainte déposée,\nnon traitée"),
        ("ASSIGNED\n(Assignée)",      "Inspecteur\nattribué"),
        ("IN_PROGRESS\n(En cours)",   "Enquête\ndémarrée"),
        ("UNDER_INVESTIGATION",       "Instruction\napprofondie"),
        ("MEDIATION\n(Médiation)",    "Séance de\nmédiation planifiée"),
        ("RESOLVED\n(Résolue)",       "Accord signé,\ndossier clos"),
    ], bg_hex="DC2626")

    h2("Statuts de sortie alternatifs")
    tbl(
        ["Statut", "Code", "Déclencheur", "Suite"],
        [
            ["Clôturé",           "CLOSED",    "Désistement plaignant",           "Archivage GED"],
            ["Escaladé",          "ESCALATED", "Non-comparution × 3 ou complexité","Transmission hiérarchie"],
            ["Judiciaire",        "JUDICIAL",  "Médiation échouée ou accord rompu","Procédure Tribunal du Travail"],
        ],
        hdr_bg="DC2626",
    )

    h2("Types de plaintes")
    badge_row(["UNPAID_SALARY", "TERMINATION", "HARASSMENT", "ACCIDENT", "OVERTIME",
               "LEAVE", "CONTRACT", "DISCRIMINATION", "WORKING_CONDITIONS", "OTHER"])

    h2("Niveaux de priorité")
    tbl(
        ["Priorité", "Code", "Délai traitement cible", "Escalade auto si dépassé"],
        [
            ["Faible",   "LOW",    "30 jours",  "Non"],
            ["Normale",  "MEDIUM", "15 jours",  "Chef d'inspection"],
            ["Haute",    "HIGH",   "7 jours",   "Directeur Régional"],
            ["Urgente",  "URGENT", "48 heures", "Directeur Général"],
        ],
        hdr_bg="DC2626",
    )

    h2("Modèles")
    tbl(
        ["Modèle", "Champs clés", "Rôle"],
        [
            ["Complaint",              "complainant, enterprise, complaint_type, status, priority, assigned_inspector", "Dossier central"],
            ["ComplaintDocument",      "complaint, doc_type (CONTRACT/PAYSLIP/PHOTO/AUDIO/VIDEO/LETTER/OTHER), file",  "Pièces jointes"],
            ["ComplaintComment",       "complaint, author, content, is_internal",                                       "Commentaires internes/externes"],
            ["ComplaintStatusHistory", "complaint, old_status, new_status, changed_by, changed_at, note",              "Journal des transitions"],
            ["ComplaintNotification",  "complaint, recipient, notification_type, is_read, sent_at",                    "Notifications liées au dossier"],
        ],
        hdr_bg="DC2626",
    )

    h2("Endpoints clés")
    tbl(
        ["Méthode", "Endpoint", "Description"],
        [
            ["GET/POST",   "/api/complaints/",                "Lister / Créer une plainte"],
            ["GET/PUT/DEL","/api/complaints/{id}/",           "Détail / Modifier / Supprimer"],
            ["POST",       "/api/complaints/{id}/assign/",    "Assigner à un inspecteur"],
            ["POST",       "/api/complaints/{id}/update_status/", "Changer le statut"],
            ["GET",        "/api/complaints/notifications/",  "Notifications de l'utilisateur"],
        ],
        hdr_bg="DC2626",
    )

    h2("Notifications automatiques")
    tbl(
        ["Événement", "Destinataire", "Canal"],
        [
            ["Plainte déposée",    "Plaignant + Chef d'inspection du district", "Email + Push"],
            ["Assignation",        "Inspecteur assigné",                         "Email + Push"],
            ["Changement statut",  "Plaignant",                                  "Email + Push + SMS"],
            ["Commentaire ajouté", "Parties concernées",                          "Push"],
            ["Délai dépassé",      "Inspecteur + Chef d'inspection",              "Email + Push"],
        ],
        hdr_bg="F47E1C",
    )
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  6. MODULE INSPECTIONS
# ════════════════════════════════════════════════════════════════════════════
def section_inspections():
    h1("6. Module INSPECTIONS — Contrôles terrain et zones", color=VERT, bar_hex="009A44")

    body("Gère le découpage géographique de la Côte d'Ivoire en zones d'inspection, "
         "l'affectation des inspecteurs aux zones, et le suivi des contrôles sur le terrain.")

    h2("Modèles")
    tbl(
        ["Modèle", "Champs clés", "Rôle"],
        [
            ["InspectionZone",         "name, region, city, lat, long, head_inspector",                      "Zone géographique d'inspection"],
            ["Commune",                "name, zone (FK), sub_prefecture",                                     "Sous-préfecture/commune de la zone"],
            ["InspectorZoneAssignment","inspector, zone, role (HEAD/MEMBER), languages_spoken, assigned_at", "Affectation inspecteur ↔ zone"],
            ["InspectionRecord",       "inspector, enterprise, record_type, result, date, report_file",      "Compte-rendu de visite/contrôle"],
        ],
        hdr_bg="009A44",
    )

    h2("Types de contrôles et résultats")
    tbl(
        ["Type", "Code", "Résultat possible", "Action si non-conforme"],
        [
            ["Contrôle de routine",        "ROUTINE",       "COMPLIANT / MINOR_ISSUES / MAJOR_ISSUES / NON_COMPLIANT", "Mise en demeure"],
            ["Contrôle de suivi",           "FOLLOW_UP",     "COMPLIANT / MAJOR_ISSUES / NON_COMPLIANT",                "PV d'infraction"],
            ["Contrôle sur plainte",        "COMPLAINT",     "COMPLIANT / NON_COMPLIANT",                               "Escalade dossier plainte"],
            ["Contrôle inoppiné",           "SPOT_CHECK",    "COMPLIANT / MINOR_ISSUES / MAJOR_ISSUES / NON_COMPLIANT", "Rapport immédiat chef"],
        ],
        hdr_bg="009A44",
    )

    h2("Workflow d'un contrôle terrain")
    flow([
        ("Planification\n(Agenda)",       "Inspecteur planifie\nla visite dans l'agenda"),
        ("Arrivée\n(QR Scan)",            "Scan QR code\nentreprise sur mobile"),
        ("Contrôle\n(Checklist)",         "Remplissage\nchecklist mobile"),
        ("Rapport\n(Rapport visite)",     "Rédaction rapport\nconstat sur site"),
        ("Soumission\n(Sync API)",        "Envoi rapport\nvers le serveur"),
        ("Suivi\n(FOLLOW_UP)",            "Planification\ncontrôle de suivi"),
    ], bg_hex="009A44")

    h2("Endpoints clés")
    tbl(
        ["Méthode", "Endpoint", "Description"],
        [
            ["GET/POST",   "/api/inspections/zones/",           "Zones d'inspection"],
            ["GET/POST",   "/api/inspections/communes/",        "Communes"],
            ["GET/POST",   "/api/inspections/zone-assignments/","Affectations inspecteurs"],
            ["GET/POST",   "/api/inspections/records/",         "Comptes-rendus de contrôle"],
            ["GET",        "/api/inspections/records/?inspector={id}", "Contrôles d'un inspecteur"],
        ],
        hdr_bg="009A44",
    )
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  7. MODULE MEDIATIONS
# ════════════════════════════════════════════════════════════════════════════
def section_mediations():
    h1("7. Module MEDIATIONS — Médiation et conciliation", color=TEAL, bar_hex="0F766E")

    body("Module central de résolution amiable. Gère le cycle complet des séances de "
         "médiation : création, convocation multi-canal traçable, accusé de réception, "
         "gestion des non-comparutions, rédaction et signature des accords.")

    h2("7.1 Cycle de vie d'une séance")
    flow([
        ("SCHEDULED\n(Planifiée)",  "Séance créée,\nconvocations envoyées"),
        ("ONGOING\n(En cours)",     "Séance démarrée\npar le médiateur"),
        ("COMPLETED\n(Terminée)",   "Séance close avec\nou sans accord"),
        ("POSTPONED\n(Reportée)",   "Renvoi à une\nnouvelle date"),
        ("CANCELLED\n(Annulée)",    "Annulation\ndéfinitive"),
    ], bg_hex="0F766E")

    h2("7.2 Workflow de convocation — Accusé de réception traçable")
    body("Chaque participant reçoit un jeton UUID unique. Le lien dans l'email/SMS "
         "enregistre l'accusé de réception sans nécessiter de connexion.")
    flow([
        ("Non convoqué",          "Participant\najouté à la séance"),
        ("Email + SMS + Push\nenvoyés", "3 canaux\nsimultanément"),
        ("Clic lien UUID\n(public)", "acknowledged_at\nenregistré"),
        ("Confirmation\nmanuelle", "Inspecteur confirme\nen présentiel"),
        ("3 relances\nsans réponse", "Escalade auto\nhiérarchie"),
    ], bg_hex="7C3AED")

    h2("7.3 Gestion des non-comparutions")
    tbl(
        ["Action déclenchée", "Détail"],
        [
            ["employer_no_show = True",      "Champ enregistré avec timestamp et auteur (inspecteur)"],
            ["Statut plainte → ESCALATED",   "Ou JUDICIAL si open_infraction_pv=True"],
            ["compliance_score − 15",        "Score de conformité de l'entreprise décrémenté"],
            ["PV de carence généré",         "Document MediationDocument de type PV_CARENCE créé"],
            ["PV d'infraction (optionnel)",  "Document de type INFRACTION si l'option est cochée"],
            ["Notification hiérarchie",      "Chef d'Inspection + Directeur Régional notifiés"],
        ],
        hdr_bg="DC2626",
    )

    h2("7.4 Résultats de séance")
    tbl(
        ["Résultat", "Code", "Condition", "Suite"],
        [
            ["Accord trouvé",       "AGREEMENT",         "Toutes parties d'accord",  "Signature accord → SIGNED"],
            ["Accord partiel",      "PARTIAL_AGREEMENT", "Points partiels",           "Accord partiel + nouvelle séance"],
            ["Pas d'accord",        "NO_AGREEMENT",      "Désaccord total",           "Renvoi judiciaire"],
            ["Non-comparution",     "(voir §7.3)",        "Employeur absent",         "PV carence + escalade"],
        ],
        hdr_bg="0F766E",
    )

    h2("7.5 Cycle de vie d'un accord")
    flow([
        ("DRAFT\n(Brouillon)",        "Rédigé par\nle médiateur"),
        ("PENDING_SIGNATURE",         "En attente\ndes signatures"),
        ("SIGNED\n(Signé)",           "3 signatures\ncollectées"),
        ("EXECUTED\n(Exécuté)",       "Obligations\nrespectées"),
        ("BREACHED\n(Rompu)",         "Non-respect →\nrenvoi JUDICIAL"),
    ], bg_hex="166534")

    h2("7.6 Modèles")
    tbl(
        ["Modèle", "Champs clés", "Rôle"],
        [
            ["Mediation",             "complaint, mediator, session_date, status, outcome, employer_no_show, pv_carence_generated, postpone_count", "Séance principale"],
            ["MediationParticipant",  "mediation, role, acknowledgment_token, acknowledged_at, convocation_attempts, is_acknowledged", "Participant avec suivi convocation"],
            ["ConvocationAttempt",    "participant, channel (EMAIL/SMS/PUSH/MANUAL), status, sent_at, provider_message_id", "Log de chaque tentative"],
            ["Agreement",             "mediation, terms, employee_obligations, employer_obligations, execution_deadline, status, is_fully_signed", "Accord signé"],
            ["MediationMinutes",      "mediation, opening_statement, discussion_summary, agreements_reached, signed_at", "PV de séance"],
            ["MediationDocument",     "mediation, document_type (CONVOCATION/MINUTES/AGREEMENT/PV_CARENCE/INFRACTION), file", "Documents attachés"],
        ],
        hdr_bg="0F766E",
    )

    h2("7.7 Endpoints complets")
    tbl(
        ["Méthode", "Endpoint", "Description", "Auth"],
        [
            ["GET/POST",  "/api/mediations/sessions/",                         "Lister / Créer séance",             "INSPECTEUR+"],
            ["GET/PUT",   "/api/mediations/sessions/{id}/",                    "Détail / Modifier",                  "INSPECTEUR+"],
            ["POST",      "/api/mediations/sessions/{id}/convoke/",            "Envoyer convocations",               "INSPECTEUR+"],
            ["GET",       "/api/mediations/sessions/acknowledge/{token}/",     "Accuser réception (lien)",           "PUBLIC"],
            ["POST",      "/api/mediations/sessions/{id}/manual_acknowledge/", "Confirmer manuellement",             "INSPECTEUR+"],
            ["POST",      "/api/mediations/sessions/{id}/relance/",            "Relancer non-accusés",               "INSPECTEUR+"],
            ["POST",      "/api/mediations/sessions/{id}/report_no_show/",     "Constater non-comparution",          "INSPECTEUR+"],
            ["POST",      "/api/mediations/sessions/{id}/postpone/",           "Reporter la séance",                 "INSPECTEUR+"],
            ["POST",      "/api/mediations/sessions/{id}/start/",              "Démarrer la séance",                 "INSPECTEUR+"],
            ["POST",      "/api/mediations/sessions/{id}/complete/",           "Clôturer la séance",                 "INSPECTEUR+"],
            ["POST",      "/api/mediations/sessions/{id}/generate_minutes/",   "Générer PV de séance",               "INSPECTEUR+"],
            ["GET",       "/api/mediations/sessions/no_shows/",                "Séances avec non-comparution",       "INSPECTEUR+"],
            ["GET",       "/api/mediations/sessions/pending_acknowledgment/",  "Séances sans accusé",                "INSPECTEUR+"],
            ["POST",      "/api/mediations/agreements/{id}/sign/",             "Signer l'accord",                    "PARTIES"],
            ["POST",      "/api/mediations/agreements/{id}/mark_breached/",    "Marquer accord rompu",               "INSPECTEUR+"],
        ],
        hdr_bg="0F766E",
    )

    h2("7.8 Rotation des canaux lors des relances")
    tbl(
        ["Relance n°", "Délai depuis la convocation initiale", "Canal utilisé"],
        [
            ["1ère relance",    "48 heures",   "EMAIL"],
            ["2ème relance",    "+24 heures",  "SMS"],
            ["3ème relance",    "+24 heures",  "PUSH mobile"],
            ["Après 3 relances","Immédiat",    "Escalade hiérarchie automatique (Celery)"],
        ],
        hdr_bg="7C3AED",
    )
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  8. MODULE HIERARCHY
# ════════════════════════════════════════════════════════════════════════════
def section_hierarchy():
    h1("8. Module HIERARCHY — Escalade et approbations", color=VIOLET, bar_hex="7C3AED")

    body("Gère les circuits de validation hiérarchique, l'escalade des dossiers complexes, "
         "les délégations temporaires et la traçabilité des réaffectations.")

    h2("Modèles")
    tbl(
        ["Modèle", "Champs clés", "Rôle"],
        [
            ["WorkflowApproval",    "complaint/mediation (FK), level (INSPECTOR/CHIEF/REGIONAL/NATIONAL), status, approver, decision_note", "Circuit d'approbation"],
            ["Escalation",          "complaint (FK), from_user, to_user, reason (DELAY/COMPLEXITY/CONFLICT/MANUAL/OTHER), notes", "Escalade de dossier"],
            ["Delegation",          "delegator, delegate, start_date, end_date, scope, is_active", "Délégation temporaire de tâches"],
            ["ReassignmentHistory", "complaint, from_inspector, to_inspector, reason, reassigned_by, reassigned_at", "Historique réaffectations"],
        ],
        hdr_bg="7C3AED",
    )

    h2("Workflow d'approbation")
    flow([
        ("PENDING\n(En attente)",   "Soumis à\napprobation"),
        ("INSPECTOR\n(Inspecteur)", "Validation niveau\ninspecteur"),
        ("CHIEF\n(Chef inspect.)",  "Validation\nchef inspection"),
        ("REGIONAL\n(Dir. Rég.)",   "Validation\ndirecteur régional"),
        ("NATIONAL\n(Dir. Gén.)",   "Validation\ndirecteur général"),
        ("APPROVED / REJECTED",     "Décision\nfinale"),
    ], bg_hex="7C3AED")

    h2("Raisons d'escalade")
    badge_row(["DELAY — Délai dépassé", "COMPLEXITY — Dossier complexe",
               "CONFLICT — Conflit d'intérêt", "MANUAL — Décision manuelle", "OTHER"])

    h2("Processus de délégation")
    for step in [
        "1. Un inspecteur peut déléguer ses dossiers temporairement (congés, mission).",
        "2. La délégation a une date de début et de fin, et un périmètre (tout ou partie).",
        "3. Le délégué reçoit automatiquement les notifications des dossiers délégués.",
        "4. À l'expiration de la délégation, les dossiers retournent automatiquement à l'inspecteur initial.",
    ]:
        bul(step)
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  9. MODULE JUDICIAL
# ════════════════════════════════════════════════════════════════════════════
def section_judicial():
    h1("9. Module JUDICIAL — Procédures judiciaires", color=ROUGE, bar_hex="DC2626")

    body("Gère les dossiers transmis aux juridictions du travail : Tribunal du Travail "
         "d'Abidjan et tribunaux régionaux compétents, les audiences et les décisions.")

    h2("Workflow d'une procédure judiciaire")
    flow([
        ("PREPARATION",     "Constitution\ndu dossier"),
        ("SUBMITTED",       "Transmis au\ngreffe"),
        ("UNDER_REVIEW",    "Examen par\nle tribunal"),
        ("HEARING_SCHEDULED","Audience\nplanifiée"),
        ("AWAITING_DECISION","En attente\nde décision"),
        ("DECIDED",         "Décision\nrendue"),
        ("APPEAL",          "Appel déposé\n(optionnel)"),
        ("CLOSED",          "Procédure\ncloturée"),
    ], bg_hex="DC2626")

    h2("Types de procédures")
    badge_row(["LABOR_DISPUTE", "WRONGFUL_TERMINATION", "DISCRIMINATION", "HARASSMENT",
               "UNPAID_WAGES", "WORKPLACE_ACCIDENT", "OTHER"])

    h2("Modèles")
    tbl(
        ["Modèle", "Champs clés", "Rôle"],
        [
            ["JudicialProcedure", "complaint (FK), procedure_type, status, tribunal, filing_date, case_reference",  "Dossier judiciaire"],
            ["Hearing",           "procedure (FK), hearing_date, location, judge, outcome, postponed",               "Audience"],
            ["JudicialDecision",  "procedure (FK), decision_date, decision_text, in_favor_of, appeal_deadline",      "Décision + appel"],
        ],
        hdr_bg="DC2626",
    )

    h2("Cas déclencheurs")
    tbl(
        ["Déclencheur", "Action"],
        [
            ["Médiation sans accord (NO_AGREEMENT)",     "Statut plainte → JUDICIAL, création JudicialProcedure"],
            ["Accord de médiation rompu (BREACHED)",     "Statut plainte → JUDICIAL via mark_breached()"],
            ["Non-comparution × 3 + escalade",          "Directeur peut décider renvoi judiciaire"],
            ["Infraction grave (PV d'infraction)",       "Ouverture procédure distincte"],
        ],
        hdr_bg="DC2626",
    )
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  10. MODULE DOMESTIC
# ════════════════════════════════════════════════════════════════════════════
def section_domestic():
    h1("10. Module DOMESTIC — Travailleurs domestiques", color=ORANGE, bar_hex="F47E1C")

    body("Module spécialisé pour les travailleurs domestiques (employés de maison) "
         "et leurs employeurs particuliers. Couvre le contrat, le pointage quotidien, "
         "les congés, les bulletins de paie et les plaintes vocales.")

    h2("Modèles")
    tbl(
        ["Modèle", "Champs clés", "Rôle"],
        [
            ["DomesticWorker",   "user, specialization (HOUSEKEEPER/COOK/NANNY/DRIVER/GARDENER/SECURITY/CARETAKER/GENERAL), assigned_inspector", "Profil travailleur domestique"],
            ["DomesticEmployer", "user, employer_type (INDIVIDUAL/FAMILY/COMPANY)",                                                               "Employeur particulier"],
            ["DomesticContract", "worker, employer, start_date, end_date, salary, status (DRAFT→ACTIVE→TERMINATED)",                              "Contrat de travail"],
            ["TimeTracking",     "worker, date, check_in, check_out, hours_worked, validated_by_employer",                                        "Pointage quotidien"],
            ["MonthlyPayslip",   "worker, month, year, base_salary, overtime, deductions, net_salary, pdf_file",                                  "Bulletin de paie mensuel"],
            ["LeaveRequest",     "worker, leave_type, start_date, end_date, status (PENDING/APPROVED/REJECTED)",                                   "Demande de congé"],
            ["OvertimeSession",  "worker, date, hours, reason, approved",                                                                          "Heures supplémentaires"],
            ["VoiceComplaint",   "worker, audio_file, transcript, status, created_at",                                                             "Plainte vocale enregistrée"],
            ["FieldVisit",       "inspector, worker, visit_date, findings, recommendations, report_file",                                          "Visite terrain inspecteur"],
        ],
        hdr_bg="F47E1C",
    )

    h2("Workflow du contrat domestique")
    flow([
        ("DRAFT\n(Brouillon)",        "Contrat rédigé\npar l'employeur"),
        ("PENDING_SIGNATURE",         "Envoyé au\ntravailleur"),
        ("ACTIVE\n(Actif)",           "Signé par\nles deux parties"),
        ("SUSPENDED\n(Suspendu)",     "Suspension\ntemporaire"),
        ("TERMINATED\n(Résilié)",     "Fin de\ncontrat"),
    ], bg_hex="F47E1C")

    h2("Workflow du pointage")
    for step in [
        "1. Le travailleur ouvre l'app mobile et enregistre son arrivée (check_in).",
        "2. En fin de journée, il enregistre son départ (check_out).",
        "3. Le système calcule automatiquement les heures travaillées.",
        "4. L'employeur valide ou conteste les heures depuis son espace.",
        "5. Les heures validées alimentent automatiquement le bulletin de paie mensuel.",
        "6. L'inspecteur assigné peut consulter le journal de pointage lors d'une visite.",
    ]:
        bul(step)

    h2("Plaintes vocales")
    body("Les travailleurs domestiques peuvent déposer une plainte par message vocal "
         "(accessible sans maîtrise de l'écrit). Le système transcrit l'audio via "
         "OpenAI Whisper et crée automatiquement une plainte structurée.")
    flow([
        ("Enregistrement\naudio (mobile)",  "Worker appuie\nbouton micro"),
        ("Upload\nAPI /voice-complaints/",  "Fichier audio\nenvoyé"),
        ("Transcription\nWhisper",          "Texte généré\nautomatiquement"),
        ("Validation\ninspecteur",          "Inspecteur vérifie\nla transcription"),
        ("Création\nplainte",               "Complaint créée\nsur le dossier"),
    ], bg_hex="F47E1C")
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  11. MODULE AI
# ════════════════════════════════════════════════════════════════════════════
def section_ai():
    h1("11. Module AI — Intelligence artificielle", color=VIOLET, bar_hex="7C3AED")

    body("Fournit des fonctionnalités d'IA au service des inspecteurs et des travailleurs : "
         "chatbot légal, analyse de documents, détection de plaintes similaires, "
         "prédiction de risques et détection d'abus.")

    h2("Modèles")
    tbl(
        ["Modèle", "Champs clés", "Rôle"],
        [
            ["ChatConversation",   "user, topic, model_used, total_tokens",                                              "Session de chat avec le chatbot légal"],
            ["ChatMessage",        "conversation, role (USER/ASSISTANT/SYSTEM), content, tokens_used",                   "Message individuel"],
            ["DocumentAnalysis",   "user, doc_type, file, status (PENDING→PROCESSING→COMPLETED/FAILED), analysis_result","Analyse OCR+IA de document"],
            ["ComplaintSimilarity","complaint_a, complaint_b, similarity_score, detected_at",                            "Détection de plaintes similaires"],
            ["RiskPrediction",     "complaint/enterprise, risk_score, risk_factors, prediction_date",                    "Prédiction de risque"],
            ["AbuseDetection",     "enterprise, abuse_type, confidence_score, evidence, flagged_at",                     "Détection d'abus systémiques"],
        ],
        hdr_bg="7C3AED",
    )

    h2("Workflow d'analyse documentaire")
    flow([
        ("Upload doc\n/api/ai/documents/",  "Contrat, fiche\npaie, lettre…"),
        ("PENDING\n(En file)",              "Tâche Celery\ncréée"),
        ("PROCESSING\n(En cours)",          "OCR + GPT-4\nanalyse le doc"),
        ("COMPLETED\n(Terminé)",            "Résultat JSON\nstocké"),
        ("Rapport\ninspecteur",             "Résumé affiché\ndans l'interface"),
    ], bg_hex="7C3AED")

    h2("Types de documents analysés")
    badge_row(["CONTRACT — Contrat de travail", "PAYSLIP — Bulletin de paie",
               "TERMINATION — Lettre de licenciement", "COMPLAINT — Courrier de plainte",
               "ID_CARD — Pièce d'identité", "OTHER"])

    h2("Chatbot légal")
    body("Le chatbot répond aux questions des travailleurs sur leurs droits selon le "
         "Code du Travail ivoirien. Il utilise GPT-4 avec un contexte juridique spécialisé "
         "et mémorise l'historique de la conversation.")
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  12. MODULE OBSERVATORY
# ════════════════════════════════════════════════════════════════════════════
def section_observatory():
    h1("12. Module OBSERVATORY — Statistiques nationales", color=VERT, bar_hex="009A44")

    body("Agrège les données de toute la plateforme pour produire des statistiques "
         "nationales consultables par la hiérarchie et le Directeur Général.")

    h2("Modèles")
    tbl(
        ["Modèle", "Champs clés", "Rôle"],
        [
            ["DailyStatistics",  "date, complaints_created, complaints_resolved, mediations_held, inspections_done", "Métriques journalières agrégées"],
            ["MonthlyReport",    "month, year, kpi_json, generated_at, generated_by",                               "Rapport mensuel consolidé"],
        ],
        hdr_bg="009A44",
    )

    h2("Endpoints statistiques")
    tbl(
        ["Endpoint", "Description"],
        [
            ["/api/observatory/national/",              "Statistiques nationales globales"],
            ["/api/observatory/trends/",                "Tendances sur une période"],
            ["/api/observatory/complaints/by-type/",    "Répartition des plaintes par type"],
            ["/api/observatory/complaints/by-region/",  "Répartition par région"],
            ["/api/observatory/enterprises/top/",       "Entreprises avec le plus de plaintes"],
            ["/api/observatory/metrics/",               "KPIs de performance de l'Inspection"],
            ["/api/observatory/dashboard/executive/",   "Tableau de bord exécutif DGT"],
        ],
        hdr_bg="009A44",
    )
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  13. MODULE NOTIFICATIONS
# ════════════════════════════════════════════════════════════════════════════
def section_notifications():
    h1("13. Module NOTIFICATIONS — Emails et SMS", color=BLEU, bar_hex="1D4ED8")

    body("Centralise la gestion des templates de notification et le suivi des envois "
         "email et SMS. Permet aux administrateurs de personnaliser les messages.")

    h2("Modèles")
    tbl(
        ["Modèle", "Champs clés", "Rôle"],
        [
            ["EmailTemplate", "name, template_type (10 types), subject_template, body_template, is_active",         "Modèle email HTML"],
            ["EmailLog",      "recipient, template, status (PENDING/SENT/FAILED/BOUNCED), opened_at, clicked_at",   "Journal des envois email"],
            ["SMSTemplate",   "name, template_type, body_template, is_active",                                       "Modèle SMS"],
            ["SMSLog",        "recipient_phone, template, status, provider_response, sent_at",                       "Journal des envois SMS"],
        ],
        hdr_bg="1D4ED8",
    )

    h2("10 types de templates email")
    badge_row([
        "COMPLAINT_CREATED", "ASSIGNED", "STATUS_UPDATE", "MEDIATION_INVITATION",
        "HEARING_NOTIFICATION", "DECISION_NOTIFICATION", "CONTRACT_SIGNED",
        "PAYSLIP_GENERATED", "LEAVE_APPROVED", "CUSTOM",
    ])

    h2("Fournisseurs SMS Côte d'Ivoire")
    tbl(
        ["Fournisseur", "Paramètre settings.py", "Indicatif", "Fallback"],
        [
            ["Orange CI",    "SMS_GATEWAY_URL + SMS_GATEWAY_KEY", "+225", "MTN CI"],
            ["MTN CI",       "SMS_MTN_URL + SMS_MTN_KEY",         "+225", "Moov Africa"],
            ["Moov Africa",  "SMS_MOOV_URL + SMS_MOOV_KEY",       "+225", "Simulation (log)"],
        ],
        hdr_bg="1D4ED8",
    )
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  14. MODULE GED
# ════════════════════════════════════════════════════════════════════════════
def section_ged():
    h1("14. Module GED — Gestion Électronique de Documents", color=GRIS, bar_hex="6B7280")

    body("Système d'archivage numérique avec versioning, contrôle d'accès, "
         "et règles de rétention. Centralise tous les documents produits par la plateforme.")

    h2("Modèles")
    tbl(
        ["Modèle", "Champs clés", "Rôle"],
        [
            ["DocumentCategory", "name, parent (FK, hiérarchique), retention_years, description",           "Catégorie avec rétention"],
            ["Document",         "title, category, status (DRAFT/ACTIVE/ARCHIVED/DELETED), confidentiality_level (1-5), linked_complaint, linked_enterprise", "Document principal"],
            ["DocumentVersion",  "document (FK), version_number, file, created_by, change_notes",           "Version du document"],
            ["DocumentAccess",   "document (FK), user, action (VIEW/DOWNLOAD/EDIT), accessed_at, ip_address","Journal d'accès"],
            ["Archive",          "documents (M2M), archived_by, archive_date, location_code",               "Archive physique/numérique"],
        ],
        hdr_bg="6B7280",
    )

    h2("Niveaux de confidentialité")
    tbl(
        ["Niveau", "Valeur", "Accès"],
        [
            ["Public",          "1", "Tout utilisateur connecté"],
            ["Interne",         "2", "Agents de l'Inspection du Travail"],
            ["Confidentiel",    "3", "Inspecteur responsable du dossier"],
            ["Très confidentiel","4","Chef d'Inspection et supérieurs"],
            ["Secret",          "5", "Directeur Général uniquement"],
        ],
        hdr_bg="6B7280",
    )
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  15. MODULE ADMINISTRATION
# ════════════════════════════════════════════════════════════════════════════
def section_administration():
    h1("15. Module ADMINISTRATION — Configuration système", color=NOIR, bar_hex="111827")

    body("Gère la configuration globale de la plateforme, les logs d'audit, "
         "les sauvegardes, et la gestion des permissions et rôles.")

    h2("Modèles")
    tbl(
        ["Modèle", "Champs clés", "Rôle"],
        [
            ["SystemConfiguration", "key, value, category, is_public",                                         "Paramètres globaux de l'application"],
            ["AuditLog",            "user, action, model, object_id, ip_address, timestamp, payload",          "Journal d'audit complet"],
            ["BackupLog",           "backup_type, status, size_mb, storage_path, created_at",                   "Historique des sauvegardes"],
            ["MaintenanceMode",     "is_active, message, started_at, planned_end",                              "Mode maintenance"],
            ["Permission",          "name, code, description, module",                                          "Permission atomique"],
            ["Role",                "name, code, permissions (M2M)",                                            "Rôle groupant des permissions"],
            ["RolePermission",      "role, permission, granted_at, granted_by",                                 "Association rôle ↔ permission"],
            ["UserAdmin",           "user, managed_by, notes",                                                  "Gestion admin des utilisateurs"],
        ],
        hdr_bg="111827",
    )

    h2("Rétention des données")
    tbl(
        ["Type de données", "Rétention", "Action en fin de période"],
        [
            ["Logs d'audit",             "5 ans",   "Archivage GED"],
            ["Logs email/SMS",           "2 ans",   "Suppression automatique"],
            ["Dossiers de plainte résolus","10 ans", "Archivage GED puis anonymisation"],
            ["PV de non-comparution",    "Permanent","Archivage GED permanent"],
            ["Accords signés",           "Permanent","Archivage GED permanent"],
            ["Décisions judiciaires",    "Permanent","Archivage GED permanent"],
        ],
        hdr_bg="111827",
    )
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  16. MODULE BI
# ════════════════════════════════════════════════════════════════════════════
def section_bi():
    h1("16. Module BI — Business Intelligence", color=BLEU, bar_hex="1D4ED8")

    body("Fournit des tableaux de bord personnalisables, des rapports planifiés "
         "et des KPIs pour le pilotage stratégique de l'Inspection du Travail.")

    h2("Modèles")
    tbl(
        ["Modèle", "Champs clés", "Rôle"],
        [
            ["Dashboard", "name, owner, widgets (JSON), layout (JSON), is_public, shared_with",       "Tableau de bord personnalisable"],
            ["Report",    "name, report_type (TABLE/CHART/PIVOT/CUSTOM), status, schedule_cron, query_config", "Rapport schedulé"],
            ["KPI",       "name, formula, current_value, target_value, unit, trend, category",         "Indicateur clé de performance"],
            ["DataExport","report (FK), format (CSV/XLSX/PDF), requested_by, file, status",            "Export de données"],
        ],
        hdr_bg="1D4ED8",
    )

    h2("KPIs standards de l'Inspection du Travail")
    tbl(
        ["KPI", "Description", "Cible"],
        [
            ["Taux de résolution",         "% plaintes résolues / total plaintes",           "> 70%"],
            ["Délai moyen traitement",      "Jours moyens de PENDING à RESOLVED",             "< 21 jours"],
            ["Taux de médiation réussie",   "% médiations avec accord / total médiations",   "> 60%"],
            ["Taux de non-comparution",     "% séances avec employer_no_show",               "< 10%"],
            ["Productivité inspecteur",     "Dossiers traités par inspecteur par mois",       "> 15"],
            ["Score moyen conformité",      "Moyenne compliance_score toutes entreprises",    "> 70/100"],
        ],
        hdr_bg="1D4ED8",
    )
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  17. MODULE LANDING
# ════════════════════════════════════════════════════════════════════════════
def section_landing():
    h1("17. Module LANDING — Site web public", color=VERT, bar_hex="009A44")

    body("Site vitrine public accessible sans authentification. Informe les travailleurs "
         "de leurs droits, publie les actualités du Ministère, expose les ressources légales "
         "et propose un assistant IA pour les questions fréquentes.")

    h2("Modèles")
    tbl(
        ["Modèle", "Champs clés", "Rôle"],
        [
            ["NewsArticle",       "title, category (COMMUNIQUE/REFORM/EVENT/CAMPAIGN), content, published_at",  "Articles d'actualité"],
            ["FAQ",               "question, answer, category, order",                                           "Questions fréquentes"],
            ["ResourceDocument",  "title, doc_type (CODE_TRAVAIL/GUIDE/CONVENTION/CONTRACT_MODEL/OTHER), file",  "Ressources légales téléchargeables"],
            ["InspectionOffice",  "name, region, city, address, phone, email, lat, long, head_inspector",       "Bureaux d'inspection régionaux"],
            ["ContactMessage",    "name, email, subject, message, is_read, replied_at",                          "Messages du formulaire de contact"],
            ["AIKeywordResponse", "keyword, response, category",                                                 "Réponses IA aux mots-clés"],
            ["Testimonial",       "author, content, rating, is_approved",                                        "Témoignages d'utilisateurs"],
            ["Campaign",          "title, description, start_date, end_date, banner_image",                      "Campagnes de sensibilisation"],
        ],
        hdr_bg="009A44",
    )
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  18. TÂCHES CELERY
# ════════════════════════════════════════════════════════════════════════════
def section_celery():
    h1("18. Tâches Celery automatiques", color=ORANGE, bar_hex="F47E1C")

    body("L'ensemble des tâches automatisées de la plateforme est orchestré par Celery "
         "avec Redis comme broker. La planification utilise django-celery-beat.")

    h2("Planning des tâches (CELERYBEAT_SCHEDULE)")
    tbl(
        ["Tâche", "Heure (Abidjan UTC+0)", "Fréquence", "Module", "Description"],
        [
            ["check_convocation_deadlines",          "07h00", "Quotidienne", "mediations", "Relance auto ou escalade après 48h sans réponse"],
            ["send_session_reminders",               "06h30", "Quotidienne", "mediations", "Rappel 24h avant séance aux parties accusées"],
            ["check_agreement_execution_deadlines",  "08h00", "Quotidienne", "mediations", "Alerte accords dont le délai d'exécution est dépassé"],
            ["auto_escalate_long_pending_mediations","10h00", "Quotidienne", "mediations", "Escalade séances bloquées 30+ jours"],
            ["aggregate_daily_statistics",           "00h30", "Quotidienne", "observatory","Agrégation des métriques de la journée"],
            ["generate_monthly_report",              "01h00", "1er du mois", "observatory","Rapport mensuel consolidé"],
            ["process_pending_document_analyses",    "Continu","À la demande","ai",         "Traitement queue OCR+IA"],
            ["cleanup_expired_tokens",               "03h00", "Quotidienne", "users",      "Suppression tokens expirés (email verify, reset)"],
            ["send_pending_notifications",           "Continu","À la demande","notifications","Envoi emails/SMS en file"],
            ["auto_assign_inspectors",               "06h00", "Quotidienne", "users",      "Auto-assignation des employés sans inspecteur"],
        ],
        hdr_bg="F47E1C",
    )

    h2("Logique de check_convocation_deadlines")
    for step in [
        "1. Sélectionne les séances en statut SCHEDULED avec participants non-accusés (acknowledged_at IS NULL).",
        "2. Si 1ère convocation date de +48h et 0 relances → envoyer relance n°1 (EMAIL).",
        "3. Si dernière relance date de +24h et nb relances < 3 → envoyer relance suivante (rotation EMAIL→SMS→PUSH).",
        "4. Si 3 relances sans réponse → ComplaintNotification au Chef d'Inspection + statut plainte → ESCALATED.",
    ]:
        bul(step)

    h2("Logique de check_agreement_execution_deadlines")
    for step in [
        "1. Sélectionne les accords en statut SIGNED dont execution_deadline est dans les 3 prochains jours ou dépassée.",
        "2. Envoie un rappel EMAIL au médiateur et au plaignant.",
        "3. Si délai dépassé de plus de 7 jours et accord toujours SIGNED → alerte Chef d'Inspection.",
    ]:
        bul(step)
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  19. RÉFÉRENTIEL API COMPLET
# ════════════════════════════════════════════════════════════════════════════
def section_api():
    h1("19. Référentiel API complet", color=NOIR, bar_hex="111827")

    h2("Base URL et authentification")
    body("Base URL : https://api.e-inspection.ci/api/")
    body("Authentification : Authorization: Bearer <access_token>")
    body("Content-Type : application/json")

    sections_api = [
        ("Users",          "1D4ED8", [
            ("POST",   "/users/register/",                    "Inscription"),
            ("POST",   "/users/login/",                       "Connexion JWT"),
            ("POST",   "/users/token/refresh/",               "Rafraîchir token"),
            ("GET/PATCH","/users/me/",                        "Profil courant"),
            ("POST",   "/users/otp/setup/",                   "Activer 2FA"),
            ("POST",   "/users/otp/verify/",                  "Vérifier TOTP"),
        ]),
        ("Enterprises",    "0F766E", [
            ("GET/POST",     "/enterprises/",                 "Liste / Créer entreprise"),
            ("GET/PUT/DEL",  "/enterprises/{id}/",            "Détail / Modifier / Supprimer"),
            ("GET/POST",     "/enterprises/branches/",        "Branches"),
            ("GET/POST",     "/enterprises/documents/",       "Documents légaux"),
        ]),
        ("Complaints",     "DC2626", [
            ("GET/POST",     "/complaints/",                  "Liste / Créer plainte"),
            ("GET/PUT",      "/complaints/{id}/",             "Détail / Modifier"),
            ("POST",         "/complaints/{id}/assign/",      "Assigner inspecteur"),
            ("POST",         "/complaints/{id}/update_status/","Changer statut"),
        ]),
        ("Inspections",    "009A44", [
            ("GET/POST",     "/inspections/zones/",           "Zones d'inspection"),
            ("GET/POST",     "/inspections/records/",         "Contrôles terrain"),
            ("GET/POST",     "/inspections/zone-assignments/","Affectations"),
        ]),
        ("Mediations",     "0F766E", [
            ("GET/POST",     "/mediations/sessions/",                         "Séances"),
            ("POST",         "/mediations/sessions/{id}/convoke/",            "Convoquer"),
            ("GET",          "/mediations/sessions/acknowledge/{token}/",     "Accuser réception"),
            ("POST",         "/mediations/sessions/{id}/report_no_show/",     "Non-comparution"),
            ("POST",         "/mediations/sessions/{id}/relance/",            "Relancer"),
            ("POST",         "/mediations/agreements/{id}/sign/",             "Signer accord"),
            ("POST",         "/mediations/agreements/{id}/mark_breached/",    "Accord rompu"),
        ]),
        ("Hierarchy",      "7C3AED", [
            ("GET/POST",     "/hierarchy/approvals/",         "Approbations"),
            ("GET/POST",     "/hierarchy/escalations/",       "Escalades"),
            ("GET/POST",     "/hierarchy/delegations/",       "Délégations"),
        ]),
        ("Judicial",       "DC2626", [
            ("GET/POST",     "/judicial/procedures/",         "Procédures"),
            ("GET/POST",     "/judicial/hearings/",           "Audiences"),
            ("GET/POST",     "/judicial/decisions/",          "Décisions"),
        ]),
        ("Domestic",       "F47E1C", [
            ("GET/POST",     "/domestic/workers/",            "Travailleurs domestiques"),
            ("GET/POST",     "/domestic/contracts/",          "Contrats"),
            ("GET/POST",     "/domestic/tracking/",           "Pointage"),
            ("GET/POST",     "/domestic/payslips/",           "Bulletins de paie"),
            ("GET/POST",     "/domestic/leaves/",             "Congés"),
            ("GET/POST",     "/domestic/voice-complaints/",   "Plaintes vocales"),
            ("GET/POST",     "/domestic/field-visits/",       "Visites terrain"),
        ]),
        ("AI",             "7C3AED", [
            ("GET/POST",     "/ai/chat/",                     "Conversations chatbot"),
            ("GET/POST",     "/ai/documents/",                "Analyse documentaire"),
            ("GET",          "/ai/similarity/",               "Plaintes similaires"),
            ("GET",          "/ai/risks/",                    "Prédictions de risque"),
        ]),
    ]

    for mod_name, color, endpoints in sections_api:
        h3(mod_name)
        tbl(
            ["Méthode", "Endpoint", "Description"],
            [(m, e, d) for m, e, d in endpoints],
            hdr_bg=color,
        )
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  20. FRONTEND
# ════════════════════════════════════════════════════════════════════════════
def section_frontend():
    h1("20. Architecture Frontend (React / TypeScript)", color=BLEU, bar_hex="1D4ED8")

    body("Application SPA React 18 + TypeScript avec React Router, Zustand (state management), "
         "Axios (API) et TailwindCSS. 154+ pages organisées par rôle et module.")

    h2("Structure des routes")
    tbl(
        ["Espace", "Préfixe route", "Pages principales", "Rôles"],
        [
            ["Authentification",    "/login, /register…",        "Login, Register, ForgotPassword, ResetPassword",                     "Tous"],
            ["Accueil",             "/",                          "Home, Services, Actualites, Ressources, FAQ, Contact, Carte",        "Public"],
            ["Espace employé",      "/espace-employe/*",          "Tableau de bord, Plaintes, Médiation, Documents",                   "EMPLOYE"],
            ["Espace employeur",    "/espace-employeur/*",        "Gestion employés, Validation heures, Convocations",                 "EMPLOYEUR"],
            ["Espace inspecteur",   "/inspecteur/*",              "Dossiers, Contrôles, Médiations, Convocations, PV, Carte",         "INSPECTEUR"],
            ["Administration",      "/admin/*",                   "Tableaux de bord, Entreprises, Plaintes, Médiations, Judiciaire, BI","CHEF+"],
        ],
        hdr_bg="1D4ED8",
    )

    h2("Services API (frontend/src/services/)")
    tbl(
        ["Fichier service", "Module backend", "Fonctions clés"],
        [
            ["api.ts",             "Tous",          "Instance Axios, intercepteur refresh token JWT, gestion 401"],
            ["auth.service.ts",    "users",         "login, register, logout, refreshToken, setupOTP, verifyOTP"],
            ["complaint.service.ts","complaints",   "getComplaints, createComplaint, assignComplaint, updateStatus"],
            ["domestic.service.ts","domestic",      "getContracts, createTimeEntry, getPayslips, requestLeave"],
            ["enterprise.service.ts","enterprises", "getEnterprises, createEnterprise, getHistory"],
            ["inspection.service.ts","inspections", "getZones, createRecord, getAssignments"],
            ["judicial.service.ts","judicial",      "getProcedures, createHearing, getDecisions"],
            ["mediation.service.ts","mediations",   "getMediations, createSession, convoke, reportNoShow, signAgreement"],
            ["landing.service.ts", "landing",       "getNews, getFAQ, getResources, submitContact"],
            ["sites.service.ts",   "administration","getConfig, getAuditLogs"],
            ["user.service.ts",    "users",         "getProfile, updateProfile, changePassword"],
        ],
        hdr_bg="1D4ED8",
    )

    h2("Pages d'administration clés (frontend/src/pages/admin/)")
    tbl(
        ["Page", "Module", "Fonctionnalité"],
        [
            ["AdminMediationSeances.tsx",      "mediations",  "Tableau des séances + KPIs + badge accusé réception + bouton non-comparution"],
            ["AdminMediationProgrammation.tsx","mediations",  "Planification des séances de médiation"],
            ["AdminMediationSignature.tsx",    "mediations",  "Signature électronique des accords"],
            ["AdminMediationPV.tsx",           "mediations",  "Génération et visualisation des PV"],
            ["AdminJudicialCreation.tsx",      "judicial",    "Création de procédures judiciaires"],
            ["AdminJudicialAudiences.tsx",     "judicial",    "Gestion des audiences"],
            ["AdminJudicialDecisions.tsx",     "judicial",    "Saisie des décisions de justice"],
            ["AdminGestionHierarchique.tsx",   "hierarchy",   "Circuit d'approbation et escalades"],
            ["AdminValidation.tsx",            "hierarchy",   "Validation hiérarchique des dossiers"],
            ["AdminOrganigramme.tsx",          "hierarchy",   "Organigramme et zones d'inspection"],
            ["AdminBIKPIs.tsx",                "bi",          "KPIs et indicateurs de performance"],
            ["AdminBICartographie.tsx",        "bi",          "Carte des plaintes et contrôles"],
        ],
        hdr_bg="1D4ED8",
    )
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  21. MOBILE
# ════════════════════════════════════════════════════════════════════════════
def section_mobile():
    h1("21. Architecture Mobile (React Native / Expo)", color=VERT, bar_hex="009A44")

    body("Application React Native + Expo 50+ pour iOS et Android. "
         "Navigation par stack avec authentification biométrique/PIN et support vocal.")

    h2("Navigation principale")
    tbl(
        ["Stack", "Screens", "Rôles concernés"],
        [
            ["Auth Stack",    "LoginScreen, RegisterScreen, OTPScreen, PINScreen, BiometricScreen, VoiceLoginScreen, VoiceRegisterScreen", "Tous"],
            ["App Stack",     "DashboardScreen, ProfileScreen, EditProfileScreen",                                                         "Tous"],
            ["Complaints",    "ComplaintsScreen, ComplaintCreateScreen, ComplaintDetailScreen",                                            "EMPLOYE"],
            ["Médiation",     "MediationsScreen (convocations, no-show pour INSPECTEUR)",                                                  "EMPLOYE + INSPECTEUR"],
            ["Domestic",      "PointageScreen, CongesScreen, BulletinsScreen, ContratScreen, VoiceDeclarationScreen, MonInspecteurScreen", "EMPLOYE_MAISON"],
            ["Inspecteur",    "DossiersTerrainScreen, ControleTerrainScreen, RapportVisiteScreen, CarteScreen, AgendaScreen, QRScanScreen, CarnetAdressesScreen", "INSPECTEUR"],
            ["Employeur",     "GestionEmployeScreen, ValidationHeuresScreen, ContratEmployeurScreen, CreerContratScreen",                  "EMPLOYEUR"],
        ],
        hdr_bg="009A44",
    )

    h2("MediationsScreen — Fonctionnalités par rôle")
    tbl(
        ["Fonctionnalité", "Rôle", "API appelée"],
        [
            ["Voir mes convocations",              "EMPLOYE / EMPLOYEUR",  "GET /mediations/sessions/"],
            ["Confirmer réception convocation",    "EMPLOYEUR",            "POST /mediations/sessions/{id}/manual_acknowledge/"],
            ["Voir statut accusé de réception",    "INSPECTEUR",           "GET /mediations/sessions/pending_acknowledgment/"],
            ["Relancer l'employeur",               "INSPECTEUR",           "POST /mediations/sessions/{id}/relance/"],
            ["Constater non-comparution",          "INSPECTEUR",           "POST /mediations/sessions/{id}/report_no_show/"],
            ["Voir PV de carence généré",          "INSPECTEUR",           "GET /mediations/sessions/{id}/"],
        ],
        hdr_bg="009A44",
    )

    h2("Services mobile (mobile/src/services/)")
    tbl(
        ["Fichier", "Fonctions clés"],
        [
            ["api.ts",              "Instance Axios avec intercepteur auto-refresh JWT, SecureStorage pour tokens"],
            ["auth.service.ts",     "login, register, logout, setupOTP, setupPIN, setupBiometric"],
            ["complaint.service.ts","getComplaints, createComplaint, uploadDocument"],
            ["domestic.service.ts", "checkIn, checkOut, requestLeave, getPayslip, signContract"],
            ["inspector.service.ts","getDossiers, createReport, scanQR, getCalendar"],
            ["voice.service.ts",    "recordVoice, uploadAudio, getTranscription"],
        ],
        hdr_bg="009A44",
    )

    h2("Authentification mobile")
    flow([
        ("Saisie\nemail/mdp",     "Écran de\nconnexion"),
        ("Envoi\nJWT tokens",     "Access + Refresh\nSecureStorage"),
        ("OTP / PIN\n2FA",        "Code TOTP ou\nPIN à 6 chiffres"),
        ("Biométrie\n(Face ID…)", "Lecture\nempreinte/face"),
        ("Dashboard\n(App Stack)","Navigation\nselon rôle"),
    ], bg_hex="009A44")
    doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  SYNTHÈSE FINALE
# ════════════════════════════════════════════════════════════════════════════
def section_synthese():
    h1("Synthèse des flux inter-modules", color=ORANGE, bar_hex="F47E1C")

    body("Le schéma ci-dessous résume les interconnexions entre les 15 modules lors du "
         "traitement complet d'un litige du travail, du dépôt de la plainte à la clôture.")

    tbl(
        ["Étape", "Module(s) impliqué(s)", "Acteur principal", "Résultat"],
        [
            ["1. Dépôt de plainte",            "COMPLAINTS + NOTIFICATIONS",    "Travailleur",       "Complaint PENDING + email confirmation"],
            ["2. Assignation inspecteur",       "COMPLAINTS + HIERARCHY + USERS","Chef d'inspection","Complaint ASSIGNED + notification inspecteur"],
            ["3. Enquête terrain",              "INSPECTIONS + GED + AI",        "Inspecteur",        "InspectionRecord + DocumentAnalysis"],
            ["4. Médiation planifiée",          "MEDIATIONS + COMPLAINTS",       "Inspecteur",        "Mediation SCHEDULED + Complaint MEDIATION"],
            ["5. Convocations envoyées",        "MEDIATIONS + NOTIFICATIONS",    "Système (auto)",    "Email + SMS + Push parties + log ConvocationAttempt"],
            ["6. Accusé de réception",          "MEDIATIONS",                    "Employeur (lien)",  "acknowledged_at enregistré"],
            ["7a. Séance tenue — accord",       "MEDIATIONS + GED",             "Médiateur",         "Agreement SIGNED + PDF GED + Complaint RESOLVED"],
            ["7b. Séance tenue — sans accord",  "MEDIATIONS + JUDICIAL",        "Médiateur",         "Complaint JUDICIAL + JudicialProcedure créée"],
            ["7c. Non-comparution employeur",   "MEDIATIONS + HIERARCHY + ENTERPRISES","Inspecteur", "PV carence + score −15 + escalade hiérarchie"],
            ["8. Procédure judiciaire",         "JUDICIAL + GED + NOTIFICATIONS","Inspecteur / Juge","Audiences + Decision rendue"],
            ["9. Statistiques & Reporting",     "OBSERVATORY + BI",             "DGT (automatique)", "KPIs actualisés + rapport mensuel"],
        ],
        hdr_bg="F47E1C",
    )

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.clear()
    r = p.add_run(
        f"Document généré le {date.today().strftime('%d/%m/%Y')} — "
        "Plateforme e-Inspection du Travail — Côte d'Ivoire — Usage interne confidentiel"
    )
    r.font.size = Pt(8)
    r.italic = True
    r.font.color.rgb = GRIS
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


# ════════════════════════════════════════════════════════════════════════════
#  ASSEMBLAGE
# ════════════════════════════════════════════════════════════════════════════
cover()
sommaire()
section_architecture()
section_rbac()
section_users()
section_enterprises()
section_complaints()
section_inspections()
section_mediations()
section_hierarchy()
section_judicial()
section_domestic()
section_ai()
section_observatory()
section_notifications()
section_ged()
section_administration()
section_bi()
section_landing()
section_celery()
section_api()
section_frontend()
section_mobile()
section_synthese()

out = r"c:\Users\HP\Desktop\projets\plateforme-travail\backend\Workflows_Complets_eInspection_CI_2026.docx"
doc.save(out)
print(f"OK  Document genere : {out}")
print(f"    Pages estimees  : ~55-60 pages")
